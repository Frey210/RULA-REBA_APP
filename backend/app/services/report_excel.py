import math
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break
from PIL import Image as PillowImage

from app.services.report_content import ANGLE_LABELS, flatten, label, references


def write_excel(output_path, data, detections, archive):
    workbook = Workbook()
    workbook.remove(workbook.active)
    refs = references(data)
    workers = {row["id"]: row for row in data["workers"]}
    registry_names = {row['worker_id']: row['name'] for row in data['workers'] if row['worker_id']}
    blue, ink = "006D91", "172D38"

    def worker_name(worker_id):
        return workers.get(worker_id, {}).get("name", "Unassigned")

    def date(value):
        return datetime.strptime(value, "%Y-%m-%d %H:%M:%S UTC") if value and value != "-" else None

    def cell(sheet, row, column, value):
        target = sheet.cell(row, column)
        if isinstance(value, str):
            target.value = ILLEGAL_CHARACTERS_RE.sub("", value)
            target.data_type = "s"  # Exported user text must never become an Excel formula.
        else:
            target.value = value
        target.font = Font(name="Calibri", size=11, color=ink)
        target.alignment = Alignment(vertical="top", wrap_text=True)
        if isinstance(value, datetime):
            target.number_format = 'yyyy-mm-dd hh:mm:ss'
        elif isinstance(value, float):
            target.number_format = '0.0'
        elif isinstance(value, int):
            target.number_format = '#,##0'
        if row % 2 == 0:
            target.fill = PatternFill("solid", fgColor="F0F5F7")
        return target

    def sheet(name, headers, widths, rows, note=None, *, filters=True):
        result = workbook.create_sheet(name)
        result.sheet_view.showGridLines = False
        result.sheet_properties.pageSetUpPr.fitToPage = True
        result.sheet_properties.tabColor = blue
        result.page_setup.orientation = 'landscape'
        result.page_setup.paperSize = result.PAPERSIZE_A4
        result.page_setup.fitToWidth, result.page_setup.fitToHeight = 1, 0
        result.print_title_rows = '1:4'
        result.oddFooter.center.text = 'ERGoSIGHT | Page &P of &N'
        result.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        cell(result, 1, 1, f"ERGoSIGHT - {name}").font = Font(name="Calibri", size=19, bold=True, color=blue)
        result.row_dimensions[1].height = 32
        result.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
        cell(result, 2, 1, note or f"{data['session']['code']} | All times UTC | Blank numeric cells mean not measured.")
        result.row_dimensions[2].height = 32
        for column, (heading, width) in enumerate(zip(headers, widths), 1):
            result.column_dimensions[get_column_letter(column)].width = width
            target = cell(result, 4, column, heading)
            target.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            target.fill = PatternFill("solid", fgColor=blue)
        result.row_dimensions[4].height = 32
        for row_number, values in enumerate(rows, 5):
            lines = 1
            for column, value in enumerate(values, 1):
                cell(result, row_number, column, value)
                if isinstance(value, str):
                    lines = max(lines, sum(max(1, math.ceil(len(line) / max(1, widths[column - 1] - 3))) for line in value.split('\n')))
            result.row_dimensions[row_number].height = min(360, max(30, lines * 15 + 8))
        result.freeze_panes = 'B5'
        if filters and result.max_row >= 5:
            result.auto_filter.ref = f'A4:{get_column_letter(len(headers))}{result.max_row}'
        result.print_options.horizontalCentered = True
        result.print_area = f'A1:{get_column_letter(len(headers))}{max(5, result.max_row)}'
        return result

    summary = [
        ['Session', data['session']['code']], ['Status', label(data['session']['status'])],
        ['Operator', data.get('operator_full_name') or data['generated_by']], ['Generated (UTC)', date(data['generated_at'])],
        ['Started (UTC)', date(data['session']['started_at'])], ['Stopped (UTC)', date(data['session']['stopped_at'])],
        ['Worker tracks (not unique people)', len(data['workers'])], ['Events', len(data['events'])],
        ['Reviewed assessments', sum(row['status'] == 'reviewed' for row in data['assessments'])],
        ['Available photos', sum('archive_file' in row for row in data['photos'])],
        ['High-risk exposure (s, overlap merged)', data['high_risk_duration_ms'] / 1000],
        ['Application version', data['app_version']],
        ['Assessment limitation', 'Single-camera 2D estimates are provisional. Force, activity, coupling and wrist twist require operator review. Registry and legacy photos are not verified posture evidence.'],
    ]
    notes = data['session'].get('notes') or ''
    summary.extend([['Session notes' if offset == 0 else 'Notes (continued)', notes[offset:offset + 600]] for offset in range(0, len(notes), 600)])
    sheet('Summary', ['Item', 'Value'], [42, 110], summary, filters=False)
    sheet('Workers', ['Ref', 'Name', 'Employee', 'Department', 'Position', 'Camera', 'Track', 'Identity', 'Observed (s)', 'High risk (s)', 'RULA avg', 'RULA peak', 'REBA avg', 'REBA peak'],
        [12, 28, 20, 24, 24, 16, 20, 22, 18, 18, 16, 16, 16, 16],
        [[refs[w['id']], w['name'], w['employee_number'], w['department'], w['position'], refs.get(w['camera_node_id']),
          w['edge_worker_id'], label(w['identity_status']), w['observed_duration_ms'] / 1000, w['high_risk_duration_ms'] / 1000,
          w['rula']['provisional_average'], w['rula']['provisional_peak'], w['reba']['provisional_average'], w['reba']['provisional_peak']] for w in data['workers']])
    sheet('Events', ['Ref', 'Worker', 'Camera', 'Event', 'Status', 'Severity', 'Started (UTC)', 'Ended (UTC)', 'Duration (s)', 'Method', 'Provisional score', 'Risk'],
        [12, 28, 16, 30, 20, 18, 25, 25, 18, 14, 22, 20],
        [[refs[e['id']], worker_name(e['session_worker_id']), refs.get(e['camera_node_id']), label(e['event_type']), e['status'], e['severity'],
          date(e['started_at']), date(e['ended_at']), e['duration_ms'] / 1000 if e['duration_ms'] is not None else None,
          e['score_type'], e['score'], e['risk_level']] for e in data['events']])
    sheet('Assessments', ['Ref', 'Event', 'Worker', 'Method', 'Status', 'Score', 'Risk', 'Original provisional', 'Group A', 'Group B', 'Calculated (UTC)'],
        [12, 12, 28, 16, 20, 14, 24, 24, 16, 16, 25],
        [[refs[a['id']], refs.get(a['event_id']), worker_name(a['session_worker_id']), a['type'].upper(), a['status'], a['score'],
          a['risk_level'], a['provisional_score'], a['score_a'], a['score_b'], date(a['calculated_at'])] for a in data['assessments']])
    details = []
    for assessment in data['assessments']:
        values = list(flatten(assessment['breakdown'] or {})) + [('Review notes', assessment['review_notes'])]
        for key, value in values:
            chunks = [value[offset:offset + 600] for offset in range(0, len(value), 600)] if isinstance(value, str) and value else [value]
            details.extend([[refs[assessment['id']], key, index + 1, part] for index, part in enumerate(chunks)])
    sheet('Review Details', ['Assessment', 'Parameter', 'Part', 'Recorded value'], [18, 48, 12, 100], details)
    sheet('Measurements', ['Detection ID', 'Camera', 'Track', 'Frame', 'Captured (UTC)', *[f'{value} (deg)' for value in ANGLE_LABELS]],
        [42, 16, 20, 16, 25, *([23] * len(ANGLE_LABELS))],
        [[d.id, refs.get(d.camera_node_id), d.edge_worker_id, d.frame_id, d.observed_at.replace(tzinfo=None),
          *[(d.angles or {}).get(key) for key in ANGLE_LABELS]] for d in detections.values()],
        note='Representative source measurements used by events/photos. All raw samples are in detections.csv inside the ZIP. Blank = not measured; times UTC.')
    photos = sheet('Photos', ['Photo', 'Worker / event', 'Camera', 'Captured (UTC)', 'Source / status', 'Image'],
        [14, 32, 16, 25, 40, 56],
        [[refs[p['id']], (worker_name(p['session_worker_id']) if p.get('session_worker_id') else registry_names.get(p.get('worker_id'), 'Registry reference')) + '\n' + refs.get(p.get('event_id'), '-'),
          refs.get(p.get('camera_node_id')), date(p['captured_at']), label(p['source']) + '\n' + p.get('omission', p['status']), None] for p in data['photos']],
        note='Registry references and legacy photos are not verified posture evidence. Source images remain unchanged in the ZIP.', filters=False)
    for row_number, photo in enumerate(data['photos'], 5):
        photos.row_dimensions[row_number].height = 160
        if 'archive_file' in photo:
            with PillowImage.open(BytesIO(archive.read(photo['archive_file']))) as original:
                original.thumbnail((360, 190))
                stream = BytesIO()
                original.convert('RGB').save(stream, 'PNG')
            stream.seek(0)
            photos.add_image(Image(stream), f'F{row_number}')
    for row_number in range(7, photos.max_row + 1, 2):
        photos.row_breaks.append(Break(id=row_number - 1))
    technical = [['Report', 'Report ID', data['report_id']], ['Session', 'Session ID', data['session']['id']]]
    for group in ('cameras', 'workers', 'events', 'assessments', 'photos'):
        for record in data[group]:
            for key in ('id', 'device_id', 'worker_id', 'camera_node_id', 'run_id', 'event_id', 'detection_id', 'source_detection_id', 'last_detection_id', 'source_detection_ids', 'sha256', 'archive_file'):
                value = record.get(key)
                if value is not None:
                    for field, part in flatten({key: value}):
                        technical.append([refs[record['id']], field, part])
    sheet('References', ['Reference', 'Field', 'Value'], [20, 32, 100], technical)
    workbook.save(output_path)
