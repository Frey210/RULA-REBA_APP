import hashlib
import json
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from sqlalchemy import select

from app.core.security import create_access_token
from app.models.assessment import Assessment
from app.models.ergonomic_event import ErgonomicEvent
from app.models.report import Report
from app.models.snapshot import Snapshot
from app.models.session_worker import SessionWorker
from app.models.worker import Worker
from app.models.worker_enrollment_image import WorkerEnrollmentImage
from app.services.detection_ingest import persist_detection_event
from app.services.report_generation import _angle_rows, _csv
from app.services.snapshot_capture import capture_pending_evidence
from test_websocket_hub import recording, sample  # noqa: F401


def test_full_archive_is_private_uncapped_and_immutable(client, recording, db_session):
    event = sample(recording)
    event.clock_source = "backend_session_clock"
    persist_detection_event(db_session, event)
    capture_pending_evidence()
    original = db_session.scalar(select(ErgonomicEvent).where(ErgonomicEvent.event_type == "worker_observed"))
    for index in range(25):
        extra = ErgonomicEvent(session_id=original.session_id, session_worker_id=original.session_worker_id,
            camera_node_id=original.camera_node_id, event_type="worker_observed", severity="info", status="resolved",
            started_at=datetime.now(UTC), source_detection_id=original.source_detection_id, duration_ms=1000)
        db_session.add(extra)
        db_session.flush()
        db_session.add(Assessment(session_id=original.session_id, session_worker_id=original.session_worker_id,
            ergonomic_event_id=extra.id, assessment_type="reba", assessment_status="reviewed", score=3,
            risk_level="low", breakdown={"manual_inputs": {"load_score": 1, "coupling_score": 2},
            "component_scores": {"wrist": 2}}, review_notes="<check & verify> " * 300,
            calculated_at=datetime.now(UTC)))
    recording[3].status = "review_pending"
    db_session.commit()
    headers = {"Authorization": f"Bearer {create_access_token(recording[0].id)}"}
    url = f"/api/v1/reports/sessions/{recording[3].id}"
    response = client.post(url, headers=headers)
    assert response.status_code == 200, response.text
    report = response.json()
    assert report["metadata_json"]["reviewed_assessment_count"] == 25
    assert report["metadata_json"]["event_count"] >= 26
    assert "archive_path" not in report["metadata_json"]
    assert "excel_path" not in report["metadata_json"]
    excel = client.get(report["excel_url"], headers=headers)
    assert excel.status_code == 200
    workbook = load_workbook(BytesIO(excel.content))
    assert workbook.sheetnames == ['Summary', 'Workers', 'Events', 'Assessments', 'Review Details', 'Measurements', 'Photos', 'References']
    assert workbook['Assessments'].max_row >= 29
    assert workbook['Workers'].freeze_panes == 'B5'
    assert workbook['Events'].auto_filter.ref
    assert workbook['Summary']['B8'].is_date
    assert all(cell.data_type != 'f' for sheet in workbook for row in sheet for cell in row)
    assert client.get(report["excel_url"]).status_code == 401
    content = client.get(report["archive_url"], headers=headers).content
    with ZipFile(BytesIO(content)) as archive:
        manifest = json.loads(archive.read("manifest.json"))
        assert len(manifest["assessments"]) >= 25
        assert len(manifest["events"]) >= 26
        assert any(photo.get("event_id") == original.id and photo.get("archive_file") for photo in manifest["photos"])
        assert "backend_session_clock" in archive.read("detections.csv").decode("utf-8-sig")
        assert manifest["pdf_sha256"] == hashlib.sha256(archive.read("ergonomic_report.pdf")).hexdigest()
        assert manifest['excel_sha256'] == hashlib.sha256(archive.read('ergonomic_report.xlsx')).hexdigest()
        assert archive.read('ergonomic_report.xlsx') == excel.content
        assert "credential_nonce" not in json.dumps(manifest)
        assert "_path" not in json.dumps(manifest)
        for photo in manifest["photos"]:
            if photo.get("archive_file"):
                assert hashlib.sha256(archive.read(photo["archive_file"])).hexdigest() == photo["sha256"]
    other_headers = {"Authorization": f"Bearer {create_access_token(recording[1].id)}"}
    assert client.get(report["archive_url"], headers=other_headers).status_code == 404
    assert client.get(report["excel_url"], headers=other_headers).status_code == 404
    assert client.get(report["archive_url"]).status_code == 401
    second = client.post(url, headers=headers).json()
    assert second["id"] != report["id"]
    assert client.get(report["archive_url"], headers=headers).content == content


def test_photo_integrity_omissions_and_media_boundaries(client, recording, db_session, tmp_path):
    persist_detection_event(db_session, sample(recording))
    capture_pending_evidence()
    photos = list(db_session.scalars(select(Snapshot)))
    assert photos
    for photo in photos:
        photo.metadata_json = {**photo.metadata_json, "checksum_sha256": "wrong"}
    recording[3].status = "review_pending"
    db_session.commit()
    headers = {"Authorization": f"Bearer {create_access_token(recording[0].id)}"}
    response = client.post(f"/api/v1/reports/sessions/{recording[3].id}", headers=headers)
    assert response.status_code == 200
    report = response.json()
    assert report["metadata_json"]["photo_count"] == 0
    assert report["metadata_json"]["omitted_photo_count"] == len(photos)
    with ZipFile(BytesIO(client.get(report["archive_url"], headers=headers).content)) as archive:
        assert not any(name.startswith("photos/") for name in archive.namelist())
        assert all("omission" in p for p in json.loads(archive.read("manifest.json"))["photos"])
    row = db_session.get(Report, report["id"])
    row.file_path = str(tmp_path.parent / "outside.pdf")
    row.metadata_json = {**row.metadata_json, "archive_path": row.file_path, "excel_path": row.file_path}
    db_session.commit()
    assert client.get(report["download_url"], headers=headers).status_code == 404
    assert client.get(report["archive_url"], headers=headers).status_code == 404
    assert client.get(report["excel_url"], headers=headers).status_code == 404


def test_report_rejects_transitional_session_states(client, recording, db_session):
    headers = {"Authorization": f"Bearer {create_access_token(recording[0].id)}"}
    for state in ("preparing", "running", "stopping"):
        recording[3].status = state
        db_session.commit()
        assert client.post(f"/api/v1/reports/sessions/{recording[3].id}", headers=headers).status_code == 409


def test_csv_does_not_execute_user_supplied_formulas():
    content = _csv([dict(note=" \t=HYPERLINK(\"bad\")", score=-3, name="normal")], ["note", "score", "name"])
    assert "' \t=HYPERLINK" in content
    assert ",-3,normal" in content


def test_measurement_table_never_calls_a_legacy_twist_score_an_angle():
    rows = dict(_angle_rows(SimpleNamespace(angles={'neck': 0, 'wrist_twist': 1})))
    assert len(rows) == 10
    assert rows['neck'] == 0
    assert rows['wrist_l'] is None
    assert 'wrist_twist' not in rows


def test_event_photo_must_match_the_measurement_source_frame():
    from app.services.report_pdf import source_photo
    event = {'id': 'e', 'source_detection_id': 'new'}
    old = {'event_id': 'e', 'detection_id': 'old', 'archive_file': 'old.jpg'}
    correct = {'event_id': 'e', 'detection_id': 'new', 'archive_file': 'new.jpg'}
    assert source_photo(event, [old, correct]) is correct
    assert source_photo(event, [old]) is None


def test_failed_generation_preserves_existing_report_and_source_files(client, recording, db_session, monkeypatch):
    from app.services.report_generation import generate_session_pdf_report
    recording[3].status = "review_pending"
    db_session.commit()
    first = generate_session_pdf_report(db_session, recording[3], recording[0])
    pdf = Path(first.file_path)
    archive = Path(first.metadata_json['archive_path'])
    original = pdf.read_bytes(), archive.read_bytes()
    def fail(*args, **kwargs):
        raise OSError('simulated disk failure')
    monkeypatch.setattr('app.services.report_generation._write_pdf', fail)
    with pytest.raises(OSError, match='disk failure'):
        generate_session_pdf_report(db_session, recording[3], recording[0])
    assert (pdf.read_bytes(), archive.read_bytes()) == original
    assert len(list(pdf.parent.parent.iterdir())) == 1


def test_deleting_session_removes_archive_and_pdf(client, recording, db_session):
    from app.services.report_generation import generate_session_pdf_report
    recording[3].status = "review_pending"
    db_session.commit()
    report = generate_session_pdf_report(db_session, recording[3], recording[0])
    paths = [Path(report.file_path), Path(report.metadata_json['archive_path']), Path(report.metadata_json['excel_path'])]
    assert all(p.is_file() for p in paths)
    response = client.delete(f'/api/v1/sessions/{recording[3].id}', headers={
        'Authorization': f'Bearer {create_access_token(recording[0].id)}'})
    assert response.status_code == 204
    assert all(not p.exists() for p in paths)


def test_registry_photos_are_labelled_as_identity_references(client, recording, db_session):
    persist_detection_event(db_session, sample(recording))
    capture_pending_evidence()
    worker = Worker(owner_user_id=recording[0].id, name='Registry Worker', department='Assembly')
    db_session.add(worker)
    db_session.flush()
    track = db_session.scalar(select(SessionWorker))
    track.worker_id = worker.id
    track.identity_status = 'confirmed'
    photo = db_session.scalar(select(Snapshot))
    reference = WorkerEnrollmentImage(worker_id=worker.id, view='front', file_path=photo.file_path,
        content_type='image/jpeg', file_size=Path(photo.file_path).stat().st_size, width=640, height=360)
    db_session.add(reference)
    recording[3].status = 'review_pending'
    db_session.commit()
    headers = {'Authorization': f'Bearer {create_access_token(recording[0].id)}'}
    response = client.post(f'/api/v1/reports/sessions/{recording[3].id}', headers=headers)
    assert response.status_code == 200
    with ZipFile(BytesIO(client.get(response.json()['archive_url'], headers=headers).content)) as archive:
        manifest = json.loads(archive.read('manifest.json'))
        exported = next(p for p in manifest['photos'] if p['id'] == reference.id)
        assert exported['source'] == 'registry_reference_not_assessment_evidence'
        assert exported['kind'] == 'identity_reference_front'
        assert exported['archive_file'] in archive.namelist()
        assert manifest['workers'][0]['name'] == 'Registry Worker'
        workbook = load_workbook(BytesIO(archive.read('ergonomic_report.xlsx')))
        assert len(workbook['Photos']._images) == len(manifest['photos'])
        reference_row = 5 + next(index for index, row in enumerate(manifest['photos']) if row['id'] == reference.id)
        assert 'Registry Worker' in workbook['Photos'].cell(reference_row, 2).value


def test_excel_user_text_remains_text_and_missing_measurements_stay_blank(client, recording, db_session):
    from app.services.report_generation import generate_session_pdf_report
    event = sample(recording, angles={'neck': 0})
    persist_detection_event(db_session, event)
    capture_pending_evidence()
    recording[3].status = 'review_pending'
    recording[3].notes = '=HYPERLINK("https://invalid.example", "not a formula")'
    db_session.commit()
    report = generate_session_pdf_report(db_session, recording[3], recording[0])
    workbook = load_workbook(report.metadata_json['excel_path'])
    note = next(workbook['Summary'].cell(row, 2) for row in range(5, workbook['Summary'].max_row + 1)
                if workbook['Summary'].cell(row, 1).value == 'Session notes')
    assert note.value == recording[3].notes and note.data_type == 's'
    assert workbook['Measurements']['F5'].value == 0
    assert workbook['Measurements']['G5'].value is None
